import csv
import io
import re
import zipfile
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path

ALIASES = {'date':['date','transactiondate','txndate','valuedate'],
           'description':['description','narration','particulars','details','transactiondetails'],
           'debit':['debit','withdrawal','withdrawals','withdrawalamount','debitamount'],
           'credit':['credit','deposit','deposits','depositamount','creditamount'],
           'amount':['amount','transactionamount'], 'type':['type','drcr','crdr']}

def money(value):
    s = str(value or '').strip().replace(',', '').replace('₹','').replace('INR','').strip()
    if s in ('', '-', '—'):
        return 0
    s = re.sub(r'\s*(?:CR|DR)$', '', s, flags=re.I)
    if s.startswith('(') and s.endswith(')'):
        s = '-'+s[1:-1]
    try:
        d = Decimal(s)
        if not d.is_finite() or abs(d) > Decimal('1000000000'):
            raise ValueError('Amount outside supported range')
        return int((d*100).quantize(Decimal('1')))
    except (InvalidOperation, ValueError):
        raise ValueError('Unrecognized amount: '+s[:30])

def parse_date(value):
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%d-%m-%Y','%d/%m/%y','%d-%m-%y','%d %b %Y','%d-%b-%Y'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError('Unrecognized date. Use YYYY-MM-DD or day/month/year.')

def normalize(header):
    value = re.sub('[^a-z]', '', str(header).lower())
    return value[:-2] if value.endswith('rs') and value != 'particulars' else value

def mapping(header):
    norm = [normalize(v) for v in header]
    return {key:next((i for i,v in enumerate(norm) if v in aliases), None)
            for key, aliases in ALIASES.items()}

def table_rows(table):
    rows, warnings = [], []
    positions = None
    for line, cells in enumerate(table, 1):
        if not cells or all(not str(v or '').strip() for v in cells):
            continue
        m = mapping(cells)
        if m['date'] is not None and m['description'] is not None:
            positions = m
            continue
        if positions is None:
            continue
        def get(k):
            i = positions[k]
            return cells[i] if i is not None and i < len(cells) else ''
        if not get('date'):
            warnings.append(f'Line {line}: missing date; not imported. Check wrapped rows or totals.')
            continue
        try:
            d = parse_date(get('date'))
            desc = str(get('description') or '').strip()
            if not desc:
                raise ValueError('Missing description')
            if positions['debit'] is not None or positions['credit'] is not None:
                debit, credit = money(get('debit')), money(get('credit'))
                if debit < 0 or credit < 0 or (debit and credit):
                    raise ValueError('Ambiguous debit/credit amounts')
                amount = credit-debit
            else:
                amount = money(get('amount'))
                typ = str(get('type')).strip().lower()
                if typ in ('dr','debit','withdrawal'):
                    amount = -abs(amount)
                elif typ in ('cr','credit','deposit'):
                    amount = abs(amount)
                elif typ:
                    raise ValueError('Unrecognized transaction type')
            if not amount:
                continue
            rows.append(dict(date=d, description=desc[:300], amount_cents=amount,
                             category='Other', essential=False))
        except ValueError as e:
            warnings.append(f'Line {line}: {e}')
    return rows, warnings

def extract(filename, content):
    ext = Path(filename).suffix.lower()
    tables = []
    if ext == '.csv':
        text = content.decode('utf-8-sig')
        tables = [list(csv.reader(io.StringIO(text)))]
    elif ext == '.docx':
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise ValueError('Invalid DOCX file')
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 50*1024*1024:
                raise ValueError('DOCX expands beyond the allowed size')
        from docx import Document
        doc = Document(io.BytesIO(content))
        tables = [[[cell.text for cell in row.cells] for row in t.rows] for t in doc.tables]
        if not tables:
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            tables = [[re.split(r'\t|\s{2,}', line) for line in lines]]
    elif ext == '.pdf':
        if not content.startswith(b'%PDF'):
            raise ValueError('Invalid PDF file')
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            if len(pdf.pages) > 100:
                raise ValueError('Maximum 100 pages per statement')
            for page in pdf.pages:
                extracted = page.extract_tables()
                if extracted:
                    tables.extend(extracted)
                else:
                    text = page.extract_text(layout=True) or ''
                    tables.append([re.split(r'\s{2,}', line.strip()) for line in text.splitlines() if line.strip()])
    else:
        raise ValueError('Use PDF, DOCX or CSV. Save legacy .doc files as .docx first.')
    rows, warnings = [], []
    for t in tables:
        r, w = table_rows(t)
        rows.extend(r)
        warnings.extend(w)
    if len(rows) > 5000:
        raise ValueError('Maximum 5,000 transactions per upload')
    if not rows:
        warnings.append('No recognizable transaction table. Scanned PDFs need OCR first. Add rows in review or use the sample CSV layout.')
    return rows, warnings[:100]
